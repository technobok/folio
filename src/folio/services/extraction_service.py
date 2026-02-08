"""Text extraction from binary documents."""

import email
import io
import logging

log = logging.getLogger(__name__)


def extract_text(content_bytes: bytes, mime_type: str) -> str:
    """Extract plain text from binary content based on MIME type.

    Returns extracted text or empty string on failure.
    """
    extractors = {
        "application/pdf": _extract_pdf,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_docx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": _extract_xlsx,
        "message/rfc822": _extract_email,
    }
    extractor = extractors.get(mime_type)
    if not extractor:
        return ""
    try:
        return extractor(content_bytes)
    except Exception:
        log.exception("Text extraction failed for %s", mime_type)
        return ""


def _extract_pdf(content_bytes: bytes) -> str:
    import pymupdf

    parts = []
    with pymupdf.open(stream=content_bytes, filetype="pdf") as doc:
        for page in doc:
            text = page.get_text()
            if text:
                parts.append(text)
    return "\n".join(parts)


def _extract_docx(content_bytes: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(content_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _extract_xlsx(content_bytes: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" ".join(cells))
    wb.close()
    return "\n".join(parts)


def _extract_email(content_bytes: bytes) -> str:
    msg = email.message_from_bytes(content_bytes)
    parts = []
    subject = msg.get("Subject", "")
    if subject:
        parts.append(f"Subject: {subject}")
    from_addr = msg.get("From", "")
    if from_addr:
        parts.append(f"From: {from_addr}")

    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True)
                if payload:
                    parts.append(payload.decode("utf-8", errors="replace"))
    else:
        payload = msg.get_payload(decode=True)
        if payload:
            parts.append(payload.decode("utf-8", errors="replace"))

    return "\n".join(parts)
